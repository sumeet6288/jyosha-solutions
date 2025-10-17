import io
from pypdf import PdfReader
from docx import Document
import openpyxl
from typing import Optional
import logging

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Process various document types and extract text content"""
    
    @staticmethod
    def process_pdf(file_content: bytes) -> str:
        """Extract text from PDF file"""
        try:
            pdf_file = io.BytesIO(file_content)
            reader = PdfReader(pdf_file)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}")
            raise Exception(f"Failed to process PDF: {str(e)}")
    
    @staticmethod
    def process_docx(file_content: bytes) -> str:
        """Extract text from DOCX file"""
        try:
            doc_file = io.BytesIO(file_content)
            doc = Document(doc_file)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}")
            raise Exception(f"Failed to process DOCX: {str(e)}")
    
    @staticmethod
    def process_txt(file_content: bytes) -> str:
        """Extract text from TXT file"""
        try:
            return file_content.decode('utf-8', errors='ignore').strip()
        except Exception as e:
            logger.error(f"Error processing TXT: {str(e)}")
            raise Exception(f"Failed to process TXT: {str(e)}")
    
    @staticmethod
    def process_xlsx(file_content: bytes) -> str:
        """Extract text from XLSX file"""
        try:
            xlsx_file = io.BytesIO(file_content)
            workbook = openpyxl.load_workbook(xlsx_file)
            text = ""
            
            for sheet in workbook.worksheets:
                text += f"Sheet: {sheet.title}\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
                text += "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"Error processing XLSX: {str(e)}")
            raise Exception(f"Failed to process XLSX: {str(e)}")
    
    @staticmethod
    def process_csv(file_content: bytes) -> str:
        """Extract text from CSV file"""
        try:
            return file_content.decode('utf-8', errors='ignore').strip()
        except Exception as e:
            logger.error(f"Error processing CSV: {str(e)}")
            raise Exception(f"Failed to process CSV: {str(e)}")
    
    @staticmethod
    def process_file(filename: str, file_content: bytes) -> str:
        """Process file based on extension"""
        extension = filename.lower().split('.')[-1]
        
        processors = {
            'pdf': DocumentProcessor.process_pdf,
            'docx': DocumentProcessor.process_docx,
            'doc': DocumentProcessor.process_docx,
            'txt': DocumentProcessor.process_txt,
            'xlsx': DocumentProcessor.process_xlsx,
            'xls': DocumentProcessor.process_xlsx,
            'csv': DocumentProcessor.process_csv,
        }
        
        processor = processors.get(extension)
        if not processor:
            raise Exception(f"Unsupported file type: {extension}")
        
        return processor(file_content)
    
    @staticmethod
    def format_size(size_bytes: int) -> str:
        """Format file size in human-readable format"""
        for unit in ['B', 'KB', 'MB', 'GB']:
            if size_bytes < 1024.0:
                return f"{size_bytes:.1f} {unit}"
            size_bytes /= 1024.0
        return f"{size_bytes:.1f} TB"
